import logging
import os
import json
import tempfile
import shutil
import asyncio
import aiohttp
from datetime import datetime
from typing import Dict

from telegram import Update, InlineKeyboardButton, InlineKeyboardMarkup, InputFile
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    CallbackQueryHandler,
    ConversationHandler,
    ContextTypes,
    filters
)

# ==================== LOGGING ====================
logging.basicConfig(
    format='%(asctime)s - %(name)s - %(levelname)s - %(message)s',
    level=logging.INFO
)
logger = logging.getLogger(__name__)

# ==================== TOKEN BOSHQARUVI (AVTOMATIK) ====================

def get_bot_token():
    """
    Tokenni quyidagi tartibda avtomatik qidiradi:
    1. Environment variable (BOT_TOKEN)
    2. .env fayl
    3. config.json fayl (faqat bot_token maydoni)
    4. ~/.bot_token fayl (local)
    """

    # 1. Environment variable
    token = os.getenv('BOT_TOKEN')
    if token:
        logger.info("✅ Token: Environment variable")
        return token

    # 2. .env fayl
    try:
        from dotenv import load_dotenv
        load_dotenv()
        token = os.getenv('BOT_TOKEN')
        if token:
            logger.info("✅ Token: .env fayl")
            return token
    except ImportError:
        pass

    # 3. config.json fayl (faqat bot_token)
    config_paths = [
        'config.json',
        os.path.expanduser('~/.telegram_bot_config.json'),
        '/etc/telegram_bot/config.json'
    ]

    for path in config_paths:
        if os.path.exists(path):
            try:
                with open(path, 'r') as f:
                    config = json.load(f)
                    # Faqat bot_token maydonini qidiramiz
                    token = config.get('bot_token')
                    if token:
                        logger.info(f"✅ Token: {path}")
                        return token
            except Exception:
                pass

    # 4. ~/.bot_token fayl (local development)
    token_file = os.path.expanduser('~/.bot_token')
    if os.path.exists(token_file):
        try:
            with open(token_file, 'r') as f:
                token = f.read().strip()
                if token:
                    logger.info("✅ Token: ~/.bot_token")
                    return token
        except Exception:
            pass

    # Token topilmadi
    logger.error("❌ BOT_TOKEN topilmadi!")
    logger.error("Iltimos, quyidagi usullardan birini tanlang:")
    logger.error("1. Environment variable: export BOT_TOKEN=your_token")
    logger.error("2. .env fayl: echo BOT_TOKEN=your_token > .env")
    logger.error("3. config.json: {'bot_token': 'your_token'}")
    raise ValueError("BOT_TOKEN topilmadi!")

# Tokenni olish
BOT_TOKEN = get_bot_token()

# ==================== GEMINI API KALITI (AVTOMATIK) ====================

def get_gemini_key():
    """Gemini API kalitini avtomatik qidirish"""

    # 1. Environment variable
    key = os.getenv('GEMINI_API_KEY')
    if key:
        logger.info("✅ Gemini: Environment variable")
        return key

    # 2. .env fayl
    try:
        from dotenv import load_dotenv
        load_dotenv()
        key = os.getenv('GEMINI_API_KEY')
        if key:
            logger.info("✅ Gemini: .env fayl")
            return key
    except ImportError:
        pass

    # 3. config.json
    if os.path.exists('config.json'):
        try:
            with open('config.json', 'r') as f:
                config = json.load(f)
                key = config.get('gemini_api_key', '')
                if key:
                    logger.info("✅ Gemini: config.json")
                    return key
        except Exception:
            pass

    logger.warning("⚠️ GEMINI_API_KEY topilmadi. Oddiy matn ishlatiladi.")
    return ''

GEMINI_API_KEY = get_gemini_key()
GEMINI_MODEL = "gemini-2.5-flash-lite"

# ==================== POLLINATIONS AI ====================
POLLINATIONS_BASE = "https://image.pollinations.ai/prompt"

# ==================== CONVERSATION STATES ====================
TOPIC, SLIDE_COUNT, GENERATING = range(3)

# ==================== FOYDALANUVCHI MA'LUMOTLARI ====================
user_data: Dict[int, dict] = {}

# ==================== AI FUNKSiyalari ====================

async def generate_slide_content(session: aiohttp.ClientSession, topic: str, slide_num: int, total: int) -> dict:
    """Gemini API orqali slayd matnini generatsiya qilish"""

    if not GEMINI_API_KEY:
        return fallback_content(topic, slide_num, total)

    prompt = f"""Siz professional taqdimot yaratuvchisisiz.

Mavzu: {topic}
Slayd: {slide_num}/{total}

Quyidagi formatda JAVOB BERING (faqat JSON, boshqa hech narsa):
{{
  "title": "Slayd sarlavhasi (5-8 so'z)",
  "content": "Slayd matni (3-4 ta qisqa gap, har bir yangi qatordan)",
  "image_prompt": "Rasm uchun ingliz tilida aniq tavsif (10-15 so'z)"
}}

Qoidalar:
- Sarlavha qisqa va jozibali bo'lsin
- Matn o'zbek tilida bo'lsin
- Rasm tavsifi ingliz tilida, aniq va professional bo'lsin
- Faqat JSON, boshqa hech qanday matn!"""

    try:
        url = f"https://generativelanguage.googleapis.com/v1beta/models/{GEMINI_MODEL}:generateContent?key={GEMINI_API_KEY}"

        payload = {
            "contents": [{
                "parts": [{"text": prompt}]
            }],
            "generationConfig": {
                "temperature": 0.7,
                "maxOutputTokens": 500
            }
        }

        async with session.post(url, json=payload, timeout=30) as resp:
            if resp.status == 200:
                data = await resp.json()
                text = data['candidates'][0]['content']['parts'][0]['text']

                json_str = text.strip()
                if json_str.startswith('```json'):
                    json_str = json_str[7:]
                if json_str.startswith('```'):
                    json_str = json_str[3:]
                if json_str.endswith('```'):
                    json_str = json_str[:-3]

                result = json.loads(json_str.strip())
                return result
            else:
                logger.error(f"Gemini API xatolik: {resp.status}")
                return fallback_content(topic, slide_num, total)

    except Exception as e:
        logger.error(f"Gemini xatolik: {e}")
        return fallback_content(topic, slide_num, total)

def fallback_content(topic: str, slide_num: int, total: int) -> dict:
    """Zaxira matn"""
    contents = [
        ("Kirish", f"{topic} mavzusiga umumiy kirish.\nAsosiy tushunchalar va maqsadlar."),
        ("Asosiy tushuncha", f"{topic} ning markaziy g'oyasi.\nMuhim jihatlar va xususiyatlar."),
        ("Amaliy qo'llanish", f"{topic} ni hayotda qo'llash.\nMisol va foydali maslahatlar."),
        ("Afzalliklar", f"{topic} ning asosiy yutuq va imkoniyatlari.\nNima uchun muhim?"),
        ("Xulosa", f"Asosiy xulosalar.\nKelajakdagi rivojlanish yo'nalishlari."),
    ]

    idx = min(slide_num - 1, len(contents) - 1)
    title, content = contents[idx]

    return {
        'title': title,
        'content': content,
        'image_prompt': f"{topic} professional presentation, {title}, clean modern design"
    }

async def generate_image(session: aiohttp.ClientSession, prompt: str, save_path: str) -> bool:
    """Pollinations AI orqali rasm generatsiya (BEPUL!)"""

    import urllib.parse
    encoded_prompt = urllib.parse.quote(prompt[:200])

    url = f"{POLLINATIONS_BASE}/{encoded_prompt}?width=1024&height=768&nologo=true&seed=42&enhance=true"

    try:
        async with session.get(url, timeout=60) as resp:
            if resp.status == 200:
                image_data = await resp.read()
                with open(save_path, 'wb') as f:
                    f.write(image_data)

                if os.path.getsize(save_path) > 1000:
                    return True
                else:
                    logger.warning("Rasm hajmi juda kichik")
                    return False
            else:
                logger.error(f"Pollinations xatolik: {resp.status}")
                return False
    except Exception as e:
        logger.error(f"Rasm generatsiya xatolik: {e}")
        return False

# ==================== PPTX YARATISH ====================

def create_pptx(data: dict) -> str:
    """python-pptx orqali PowerPoint taqdimot yaratish"""

    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RgbColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    temp_dir = data['temp_dir']
    pptx_path = os.path.join(temp_dir, f"{data['topic']}.pptx")

    # Title slide
    title_slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(title_slide_layout)

    # Background
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), 
        prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RgbColor(102, 126, 234)
    background.line.fill.background()

    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = data['topic']
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RgbColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.5), Inches(4.2), Inches(12.333), Inches(0.8))
    tf = sub_box.text_frame
    p = tf.paragraphs[0]
    p.text = "🤖 AI tomonidan yaratilgan taqdimot"
    p.font.size = Pt(20)
    p.font.color.rgb = RgbColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER

    # Date
    date_box = slide.shapes.add_textbox(Inches(0.5), Inches(5.5), Inches(12.333), Inches(0.5))
    tf = date_box.text_frame
    p = tf.paragraphs[0]
    p.text = datetime.now().strftime("%d.%m.%Y")
    p.font.size = Pt(14)
    p.font.color.rgb = RgbColor(200, 200, 200)
    p.alignment = PP_ALIGN.CENTER

    # Content slides
    colors = [
        (248, 249, 250), (255, 245, 245), (240, 255, 244),
        (240, 248, 255), (255, 251, 235), (245, 243, 255)
    ]

    for i, slide_data in enumerate(data['slides'], 1):
        slide_layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(slide_layout)

        # Background
        bg = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(0), Inches(0),
            prs.slide_width, prs.slide_height
        )
        bg.fill.solid()
        color = colors[i % len(colors)]
        bg.fill.fore_color.rgb = RgbColor(*color)
        bg.line.fill.background()

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.333), Inches(1))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = slide_data['title']
        p.font.size = Pt(32)
        p.font.bold = True
        p.font.color.rgb = RgbColor(44, 62, 80)
        p.alignment = PP_ALIGN.LEFT

        # Content
        content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12.333), Inches(2.5))
        tf = content_box.text_frame
        tf.word_wrap = True

        content_lines = slide_data['content'].split('\n')
        for j, line in enumerate(content_lines):
            if j == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = line
            p.font.size = Pt(18)
            p.font.color.rgb = RgbColor(52, 73, 94)
            p.space_after = Pt(12)

        # Image
        if i in data['images'] and os.path.exists(data['images'][i]):
            try:
                img_path = data['images'][i]
                slide.shapes.add_picture(
                    img_path, 
                    Inches(3.5), Inches(4.2), 
                    width=Inches(6.333)
                )
            except Exception as e:
                logger.warning(f"Rasm qo'shishda xatolik: {e}")

    prs.save(pptx_path)
    return pptx_path

# ==================== BOT HANDLERLARI ====================

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Botni ishga tushirish"""
    user_id = update.effective_user.id

    cleanup_user_data(user_id)

    user_data[user_id] = {
        'topic': '',
        'slide_count': 0,
        'slides': [],
        'images': {},
        'temp_dir': tempfile.mkdtemp(),
        'message_id': None
    }

    await update.message.reply_text(
        "🎨 *AI Taqdimot Generator*\n\n"
        "Men sizga faqat **mavzu** aytishiz kifoya — qolganini\n"
        "*Gemini AI* matn yozadi va *Pollinations AI* rasmlar chizadi!\n\n"
        "📊 Tayyor taqdimot: **PowerPoint (.pptx)** formatida!\n\n"
        "✍️ Taqdimot mavzusini kiriting:\n"
        "_Masalan: Sun'iy intellekt, O'zbekiston tarixi, Biznes reja_",
        parse_mode='Markdown'
    )
    return TOPIC

async def get_topic(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Mavzuni qabul qilish"""
    user_id = update.effective_user.id
    topic = update.message.text.strip()

    if not topic or len(topic) > 100:
        await update.message.reply_text(
            "❌ Noto'g'ri mavzu. Iltimos, 1-100 belgi orasida kiriting."
        )
        return TOPIC

    user_data[user_id]['topic'] = topic

    keyboard = [
        [InlineKeyboardButton("3 ta", callback_data="3"),
         InlineKeyboardButton("5 ta", callback_data="5"),
         InlineKeyboardButton("7 ta", callback_data="7")],
        [InlineKeyboardButton("10 ta", callback_data="10"),
         InlineKeyboardButton("15 ta", callback_data="15"),
         InlineKeyboardButton("20 ta", callback_data="20")]
    ]
    reply_markup = InlineKeyboardMarkup(keyboard)

    await update.message.reply_text(
        f"✅ Mavzu: *{topic}*\n\n"
        f"📊 Nechta slayd yaratilsin?",
        reply_markup=reply_markup,
        parse_mode='Markdown'
    )
    return SLIDE_COUNT

async def get_slide_count(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Slaydlar sonini tanlash va generatsiya boshlash"""
    query = update.callback_query
    await query.answer()

    user_id = update.effective_user.id
    count = int(query.data)
    user_data[user_id]['slide_count'] = count

    status_msg = await query.edit_message_text(
        f"🚀 *{user_data[user_id]['topic']}* taqdimoti yaratilmoqda...\n\n"
        f"⏳ Slaydlar: 0/{count}\n"
        f"🖼️ Rasmlar: 0/{count}\n\n"
        f"_Bu 1-2 daqiqa vaqt olishi mumkin..._",
        parse_mode='Markdown'
    )

    user_data[user_id]['message_id'] = status_msg.message_id

    asyncio.create_task(
        generate_presentation(update, context, user_id)
    )

    return GENERATING

async def generate_presentation(update: Update, context: ContextTypes.DEFAULT_TYPE, user_id: int):
    """Asosiy generatsiya jarayoni"""
    data = user_data[user_id]
    topic = data['topic']
    total = data['slide_count']
    chat_id = update.effective_chat.id

    async with aiohttp.ClientSession() as session:
        # 1. Slayd matnlarini generatsiya qilish
        for i in range(1, total + 1):
            try:
                await update_status(context, chat_id, data['message_id'], 
                                  f"✍️ {i}-slayd matni yozilmoqda...", i, total)

                slide_content = await generate_slide_content(session, topic, i, total)
                data['slides'].append(slide_content)

                await asyncio.sleep(0.5)

            except Exception as e:
                logger.error(f"Slayd {i} generatsiya xatolik: {e}")
                data['slides'].append(fallback_content(topic, i, total))

        # 2. Rasmlarni generatsiya qilish
        for i, slide in enumerate(data['slides'], 1):
            try:
                await update_status(context, chat_id, data['message_id'],
                                  f"🎨 {i}-slayd rasmi chizilmoqda...", i, total, image=True)

                image_path = os.path.join(data['temp_dir'], f"slide_{i}.jpg")
                success = await generate_image(session, slide['image_prompt'], image_path)

                if success:
                    data['images'][i] = image_path

                await asyncio.sleep(1)

            except Exception as e:
                logger.error(f"Rasm {i} xatolik: {e}")

        # 3. PPTX yaratish va yuborish
        await update_status(context, chat_id, data['message_id'],
                          "📦 PowerPoint yaratilmoqda...", total, total, done=True)

        await build_and_send_presentation(context, chat_id, user_id)

async def update_status(context, chat_id: int, message_id: int, 
                       status: str, current: int, total: int, image: bool = False, done: bool = False):
    """Status xabarini yangilash"""
    try:
        data = user_data.get(chat_id, {})
        topic = data.get('topic', 'Taqdimot')

        progress = "█" * current + "░" * (total - current)

        text = (
            f"🚀 *{topic}*\n\n"
            f"{progress}\n"
            f"📊 Slaydlar: {current}/{total}\n"
            f"🖼️ Rasmlar: {len(data.get('images', {}))}/{total}\n\n"
            f"⏳ {status}"
        )

        if done:
            text += "\n\n✅ *Tayyor!*"

        await context.bot.edit_message_text(
            chat_id=chat_id,
            message_id=message_id,
            text=text,
            parse_mode='Markdown'
        )
    except Exception as e:
        logger.warning(f"Status yangilash xatolik: {e}")

async def build_and_send_presentation(context, chat_id: int, user_id: int):
    """PPTX taqdimot yaratish va yuborish"""
    data = user_data[user_id]

    try:
        # PPTX yaratish
        pptx_path = create_pptx(data)

        # PPTX yuborish
        with open(pptx_path, 'rb') as f:
            await context.bot.send_document(
                chat_id=chat_id,
                document=InputFile(f, filename=f"{data['topic']}.pptx"),
                caption=f"✅ *{data['topic']}* tayyor!\n\n"
                        f"📊 Slaydlar: {data['slide_count']}\n"
                        f"🖼️ Rasmlar: {len(data['images'])}\n"
                        f"🤖 Yaratuvchi: Gemini + Pollinations AI\n\n"
                        f"📁 PowerPoint faylni yuklab oling",
                parse_mode='Markdown'
            )

        # Preview rasmini yuborish
        if data['images']:
            first_img = list(data['images'].values())[0]
            if os.path.exists(first_img):
                with open(first_img, 'rb') as f:
                    await context.bot.send_photo(
                        chat_id=chat_id,
                        photo=InputFile(f),
                        caption="🎨 Taqdimotdan namuna rasm"
                    )

        cleanup_user_data(user_id)

        await context.bot.send_message(
            chat_id=chat_id,
            text="🎉 Yangi taqdimot uchun /start ni bosing!"
        )

    except Exception as e:
        logger.error(f"Taqdimot yuborish xatolik: {e}")
        await context.bot.send_message(
            chat_id=chat_id,
            text=f"❌ Xatolik yuz berdi: {str(e)}\nIltimos, qayta urinib ko'ring."
        )
        cleanup_user_data(user_id)

def cleanup_user_data(user_id: int):
    """Tozalash"""
    if user_id in user_data:
        temp_dir = user_data[user_id].get('temp_dir', '')
        if os.path.exists(temp_dir):
            shutil.rmtree(temp_dir, ignore_errors=True)
        del user_data[user_id]

async def cancel(update: Update, context: ContextTypes.DEFAULT_TYPE) -> int:
    """Bekor qilish"""
    user_id = update.effective_user.id
    cleanup_user_data(user_id)
    await update.message.reply_text("❌ Bekor qilindi. /start")
    return ConversationHandler.END

async def help_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    """Yordam"""
    gemini_status = "✅ Ulangan" if GEMINI_API_KEY else "⚠️ API kalit yo'q (oddiy matn)"

    await update.message.reply_text(
        "🤖 *AI Taqdimot Boti*\n\n"
        "*Ishlash tartibi:*\n"
        "1. /start — Mavzu kiriting\n"
        "2. Slaydlar sonini tanlang\n"
        "3. AI avtomatik yozadi + chizadi\n"
        "4. **PowerPoint (.pptx)** faylni yuklab oling\n\n"
        f"*AI holati:*\n"
        f"📝 Matn: {gemini_status}\n"
        f"🖼️ Rasm: Pollinations AI (bepul)\n\n"
        "*Eslatma:*\n"
        "- Generatsiya 1-2 daqiqa\n"
        "- Har bir slayd unikal rasm\n"
        "- PowerPoint format",
        parse_mode='Markdown'
    )

# ==================== ASOSIY ====================

def main():
    application = Application.builder().token(BOT_TOKEN).build()

    conv_handler = ConversationHandler(
        entry_points=[CommandHandler('start', start)],
        states={
            TOPIC: [MessageHandler(filters.TEXT & ~filters.COMMAND, get_topic)],
            SLIDE_COUNT: [CallbackQueryHandler(get_slide_count, pattern='^\d+$')],
            GENERATING: [],
        },
        fallbacks=[CommandHandler('cancel', cancel)],
    )

    application.add_handler(conv_handler)
    application.add_handler(CommandHandler('help', help_cmd))

    logger.info("🚀 AI Taqdimot Boti ishga tushdi...")
    logger.info(f"📝 Gemini: {'Ulangan' if GEMINI_API_KEY else 'Mavjud emas'}")
    logger.info("🖼️ Pollinations: Bepul")
    logger.info("📊 Format: PowerPoint (.pptx)")

    application.run_polling(allowed_updates=Update.ALL_TYPES)

if __name__ == '__main__':
    main()
